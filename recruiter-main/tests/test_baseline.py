import json
import numpy as np
import pytest
from src.resume_parser import parse_document
from src.text_cleaner import clean_text
from src.vector_store import save_store, load_store
from src.jd_matcher import rank_candidates


def test_parser_formats(tmp_path):
    import pymupdf
    from docx import Document
    doc = Document()
    doc.add_paragraph('Python engineer')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'PostgreSQL'
    doc.save(tmp_path / 'resume.docx')
    assert 'PostgreSQL' in parse_document(tmp_path / 'resume.docx')
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), 'Python engineer')
    pdf.save(tmp_path / 'resume.pdf')
    pdf.close()
    assert 'Python engineer' in parse_document(tmp_path / 'resume.pdf')
    (tmp_path / 'empty.txt').write_text('   ')
    with pytest.raises(ValueError, match='No readable'):
        parse_document(tmp_path / 'empty.txt')
    (tmp_path / 'broken.pdf').write_text('broken')
    with pytest.raises(Exception):
        parse_document(tmp_path / 'broken.pdf')


def test_store_ranking_and_model_guard(tmp_path):
    records = [{'candidate': name, 'sha256': name, 'text': name} for name in ['Python', 'Accounting']]
    save_store(tmp_path, np.eye(2, dtype='float32'), records, 'test')
    index, metadata = load_store(tmp_path)
    class FakeEmbedder:
        model_name = 'test'
        def encode(self, texts):
            return np.array([[1, 0]], dtype='float32')
    embedder = FakeEmbedder()
    results = rank_candidates(index, metadata, 'Python', embedder)
    assert results[0]['candidate'] == 'Python'
    assert results[0]['similarity'] == pytest.approx(1)
    assert results[1]['similarity'] == pytest.approx(0)
    assert results[0]['matched_evidence'][0]['text'] == 'Python'
    embedder.model_name = 'wrong'
    with pytest.raises(ValueError, match='model differs'):
        rank_candidates(index, metadata, 'Python', embedder)


def test_cleaner():
    assert clean_text(' Python   developer\n\n SQL\x00 ') == 'Python developer\nSQL'
